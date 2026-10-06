"""Build 01_Tu_tuong_HCM_Dai_doan_ket.pptx.

python _build/build_deck.py            -> builds the deck (uses 12_assets/processed/* when present,
                                          otherwise _build/placeholders/* so layout can be checked early)
Speaker notes live in _build/notes.py; on-slide text lives here.
Every theoretical sentence on a slide must trace to 05_theory-verification.md.
"""
import sys
from pathlib import Path
from pptx.util import Inches, Pt
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn

sys.path.insert(0, str(Path(__file__).resolve().parent))
from deck_lib import *  # noqa
import notes as N

ROOT = Path(__file__).resolve().parents[1]
CREATED = ROOT / "12_assets" / "created"
PROC = ROOT / "12_assets" / "processed"
PH = ROOT / "_build" / "placeholders"
OUT = ROOT / "01_Tu_tuong_HCM_Dai_doan_ket.pptx"
TOTAL = 15


def img(key):
    for ext in (".png", ".jpg"):
        p = PROC / f"{key}{ext}"
        if p.exists():
            return p
    return PH / f"{key}.png"


def cre(name):
    return CREATED / name


CAP = {}  # image captions, filled from processed/captions.py if present
try:
    sys.path.insert(0, str(PROC))
    from captions import CAPTIONS as CAP  # noqa
except Exception:
    CAP = {}


def cap(key, default="Ảnh minh họa — nguồn: xem phụ lục"):
    return CAP.get(key, default)


prs = new_deck()


def band(s, h):
    """Bottom ribbon band: transparent above its curve, so nothing behind it is clipped by a hard edge."""
    pic = s.shapes.add_picture(str(cre("ribbon_band.png")), Inches(0), Inches(H - h), Inches(W), Inches(h))
    pic.name = "Dai lua"
    return pic


def strike(run_tb):
    for p in run_tb.text_frame.paragraphs:
        for r in p.runs:
            r._r.get_or_add_rPr().set("strike", "sngStrike")


# =====================================================================  S1 HERO
def s01():
    s = blank(prs)
    add_bg(s, cre("bg_burgundy.png"))
    add_img(s, cre("drum_gold_14.png"), 5.55, 0.2, 7.1, 7.1, name="Motif trong dong")
    add_img(s, img("hero"), 7.05, 0, W - 7.05, H, focus=(0.5, 0.3), name="Chan dung Ho Chi Minh")
    band(s, 1.45)
    # logo chip
    chip = rect(s, 0.6, 0.45, 2.04, 0.98, fill=C["white"], radius=14000, name="Logo chip")
    add_img_fit(s, img("fpt_logo_crop"), 0.72, 0.53, w=1.8, name="Logo Dai hoc FPT")
    text(s, 0.6, 1.72, 6.4, 0.45, "TƯ TƯỞNG HỒ CHÍ MINH", font=BODY_SEMI, size=17, color="gold2", char=5)
    text(s, 0.6, 2.18, 6.4, 0.55, "về", font=DISPLAY, size=30, italic=True, color="cream")
    text(s, 0.6, 2.56, 6.6, 2.3, ["ĐẠI ĐOÀN KẾT", "DÂN TỘC"], font=DISPLAY, size=62, bold=True,
         color="ivory", spacing=0.96, name="Title")
    text(s, 0.6, 4.9, 6.4, 0.5, "Khi khác biệt không nhất thiết dẫn đến chia rẽ", font=DISPLAY, size=22,
         italic=True, color="gold2")
    text(s, 0.6, 5.44, 6.6, 0.62, [
        [{"t": "Học phần Tư tưởng Hồ Chí Minh (HCM202)", "bold": True}],
        "Nhóm ___  ·  Lớp ___  ·  GVHD: ___  ·  Đại học FPT"], size=12.5, color="cream", spacing=1.1)
    text(s, 7.4, 7.08, 5.45, 0.3, "Hồ Chí Minh, khoảng 1946 · ảnh tư liệu, phạm vi công cộng (Wikimedia Commons)",
         size=8.5, color="cream", align="r", italic=True, name="Caption")
    transition_fade(s)
    notes(s, N.S[1])


# =====================================================================  S2 COLD OPEN
def s02():
    s = blank(prs)
    add_bg(s, cre("paper_cream.png"))
    text(s, 0.6, 0.55, 5.6, 0.4, "MỞ ĐẦU  ·  CÂU HỎI CHO CẢ LỚP", font=BODY_SEMI, size=13, color="red", char=3)
    text(s, 0.6, 1.0, 5.6, 2.2, ["ĐOÀN KẾT", "LÀ GÌ?"], font=DISPLAY, size=66, bold=True, color="burg", spacing=0.92)
    text(s, 0.6, 3.45, 5.3, 1.3, [
        [{"t": "Trả lời bằng bàn tay: ", "bold": True}, "giơ 1, 2, 3 hoặc 4 ngón tay tương ứng với A, B, C, D."],
        ], size=19, color="brown", spacing=1.1)
    text(s, 0.6, 4.72, 5.2, 1.0, "Chưa có đáp án. Cuối bài, cả lớp sẽ quay lại câu hỏi này.",
         font=DISPLAY, size=19, italic=True, color="muted")
    opts = [("A", "Mọi người cùng một quan điểm, một cách nghĩ", 1), ("B", "Có khác biệt nhưng cùng hướng tới mục tiêu chung", 2),
            ("C", "Ý kiến thiểu số phải theo đa số", 3), ("D", "Tránh tranh luận để giữ hòa khí trong tập thể", 4)]
    x0, y0, cw, ch, g = 6.3, 0.85, 3.2, 2.5, 0.25
    for i, (k, t, n) in enumerate(opts):
        cx = x0 + (i % 2) * (cw + g)
        cy = y0 + (i // 2) * (ch + g)
        card = rect(s, cx, cy, cw, ch, fill=C["white"], radius=9000, name=f"Option {k}")
        shadow(card, blur=14, dist=3, alpha=0.12)
        text(s, cx + 0.28, cy + 0.2, 1.0, 0.9, k, font=DISPLAY, size=46, bold=True, color="red")
        text(s, cx + 0.28, cy + 1.12, cw - 0.56, 1.15, t, size=19, color="ink", spacing=1.05)
        add_img_fit(s, cre(f"hand_{n}_burg.png"), cx + cw - 0.95, cy + 0.22, h=0.82)
    folio(s, 2, TOTAL, dark=False)
    transition_fade(s)
    notes(s, N.S[2])


# =====================================================================  S3 KNOWLEDGE MAP
def s03():
    s = blank(prs)
    add_bg(s, cre("paper_cream_soft.png"))
    add_img(s, img("s3bg"), 0, 0, 4.55, H, focus=(0.5, 0.5), name="Panel trai - Quoc hoi khoa I 1946")
    grad_rect(s, 0, 0, 4.55, H, [(0, C["burg"], 0.9), (0.62, C["burg"], 0.84), (0.84, C["deep"], 0.45), (1, C["deep"], 0.4)], angle=90, name="Phu toi")
    text(s, 0.55, 6.62, 3.7, 0.5, cap("s3bg"), size=8.5, italic=True, color="cream", spacing=1.0, name="Caption")
    text(s, 0.55, 0.62, 3.6, 0.4, "NHÌN LẠI CHƯƠNG 5 · MỤC II", font=BODY_SEMI, size=12, color="gold2", char=2.5)
    text(s, 0.55, 1.1, 3.85, 2.6, ["KHÔNG HỌC LẠI,", "MÀ NHÌN LẠI", "SÂU HƠN"], font=DISPLAY, size=30, bold=True,
         color="ivory", spacing=0.98)
    text(s, 0.55, 3.75, 3.6, 2.0, [
        "Bốn mảnh ghép cả lớp đã học trong bài giảng Chương 5.",
        "Hôm nay, nhóm dùng chúng để trả lời một câu hỏi khó hơn."], size=16, color="cream", spacing=1.15, para_after=8)
    cx, cy = 8.85, 3.85
    import math
    nodes = [  # numbering follows the class lecture (Chương 5, mục II): 1 Vai trò · 2 Nội dung (a, b) · 3 Hình thức tổ chức
        ("1", "VAI TRÒ", "Vấn đề có ý nghĩa chiến lược,\nquyết định thành công của cách mạng;\nmục tiêu, nhiệm vụ hàng đầu\ncủa Đảng, của dân tộc", -90),
        ("2a", "ĐOÀN KẾT\nTOÀN DÂN", "Dân là chủ thể;\nnền tảng: liên minh\ncông – nông – trí", 0),
        ("2b", "ĐIỀU KIỆN", "Truyền thống yêu nước –\nnhân nghĩa – đoàn kết ·\nkhoan dung, độ lượng ·\nniềm tin vào nhân dân", 90),
        ("3", "HÌNH THỨC\nTỔ CHỨC", "Mặt trận dân tộc\nthống nhất ·\n4 nguyên tắc hoạt động", 180),
    ]
    rx, ry = 2.95, 2.25
    pos = []
    for n, t, d, ang in nodes:
        a = math.radians(ang)
        pos.append((cx + rx * math.cos(a), cy + ry * math.sin(a)))
    for (px, py) in pos:
        line(s, cx, cy, px, py, C["gold"], width=1.25, alpha=0.7)
    core = rect(s, cx - 1.18, cy - 1.18, 2.36, 2.36, fill=C["burg"], shape=MSO_SHAPE.OVAL, name="Cau hoi trung tam")
    shadow(core, blur=20, dist=4, alpha=0.25)
    text(s, cx - 1.0, cy - 0.78, 2.0, 1.56, [
        [{"t": "ĐOÀN KẾT", "font": DISPLAY, "bold": True, "size": 19, "color": "gold2"}],
        [{"t": "có phải là", "font": DISPLAY, "italic": True, "size": 15, "color": "cream"}],
        [{"t": "GIỐNG NHAU?", "font": DISPLAY, "bold": True, "size": 19, "color": "ivory"}]], align="c", anchor="m", spacing=1.0)
    lab = {  # label box placement per node: (dx, dy, w, align)
        0: (0.42, -0.36, 3.4, "l"),
        1: (-1.15, 0.36, 2.3, "c"),
        2: (0.42, -0.40, 2.7, "l"),
        3: (-1.15, 0.36, 2.3, "c"),
    }
    for i, ((n, t, d, ang), (px, py)) in enumerate(zip(nodes, pos)):
        dot = rect(s, px - 0.27, py - 0.27, 0.54, 0.54, fill=C["gold2"], shape=MSO_SHAPE.OVAL, name=f"Node {n}")
        text(s, px - 0.27, py - 0.27, 0.54, 0.54, n, font=BODY_SEMI, size=17 if len(n) == 1 else 15, color="burg", align="c", anchor="m")
        dx, dy, w, al = lab[i]
        text(s, px + dx, py + dy, w, 1.2, [[{"t": t, "bold": True, "size": 14, "color": "burg", "char": 1.5}],
                                           [{"t": d, "size": 12, "color": "brown"}]], align=al, spacing=1.05)
    content_tags(s, ["A"], dark=False)
    folio(s, 3, TOTAL, dark=False, x=4.95)
    source_note(s, "Khung: bài giảng Chương 5 của lớp, mục II\nTS. Hà Triệu Huy, ĐH KHXH&NV – ĐHQG TP.HCM", dark=False, w=4.9, y=6.84)
    transition_fade(s)
    notes(s, N.S[3])


# =====================================================================  S4 ROLE
def s04():
    s = blank(prs)
    add_bg(s, cre("paper_cream.png"))
    add_img(s, img("s4"), 0, 0, 5.45, H, focus=(0.5, 0.45), name="Anh tu lieu")
    grad_rect(s, 0, 5.4, 5.45, 2.1, [(0, C["deep"], 0), (1, C["deep"], 0.88)], angle=90, name="Toi chan anh")
    text(s, 0.45, 6.55, 4.7, 0.7, "Lễ đài Quảng trường Ba Đình, Hà Nội, 2/9/1945\nảnh tư liệu, phạm vi công cộng (Wikimedia Commons)",
         size=10, color="cream", italic=True, name="Caption")
    x = 6.05
    text(s, x, 0.55, 5.0, 0.4, "1 · VAI TRÒ", font=BODY_SEMI, size=13, color="red", char=3)
    t1 = text(s, x, 1.02, 6.6, 0.5, "Một giải pháp tình thế, nhất thời?", font=DISPLAY, size=24, italic=True, color="muted")
    strike(t1)
    arrow(s, x + 0.2, 1.62, x + 0.2, 2.12, C["gold"], width=2)
    text(s, x, 2.2, 6.8, 1.75, ["VẤN ĐỀ CÓ Ý NGHĨA CHIẾN LƯỢC,", "QUYẾT ĐỊNH THÀNH CÔNG", "CỦA CÁCH MẠNG"],
         font=DISPLAY, size=27, bold=True, color="burg", spacing=0.98, name="Claim 1")
    text(s, x, 3.98, 6.75, 0.8, [[{"t": "và là ", "italic": True, "font": DISPLAY, "color": "muted"},
                                  {"t": "mục tiêu, nhiệm vụ hàng đầu\ncủa Đảng, của dân tộc", "bold": True, "color": "red"}]],
         size=19, spacing=1.05, name="Claim 2")
    q = rect(s, x, 4.92, 6.78, 1.66, fill=C["burg"], radius=6000, name="Quote card")
    text(s, x + 0.3, 5.1, 6.2, 1.0, N.Q["lauda"]["short"], font=DISPLAY, size=17, italic=True, color="ivory", spacing=1.05)
    text(s, x + 0.3, 6.04, 6.2, 0.42, N.Q["lauda"]["cite"].replace(" · Toàn tập", "\nToàn tập"), size=10, color="gold2", spacing=1.0)
    content_tags(s, ["A"], dark=False)
    folio(s, 4, TOTAL, dark=False, x=x)
    transition_fade(s)
    notes(s, N.S[4])


# =====================================================================  S5 AI NAM TRONG CHU DAI
def s05():
    s = blank(prs)
    add_bg(s, cre("bg_burgundy_left.png"))
    with Image.open(img("mosaic")) as _m:  # fit the collage inside its box (Phudu letters are narrow and tall)
        _mw, _mh = _m.size
    _bw, _bh = 7.0, 3.55
    _w = min(_bw, _bh * _mw / _mh)
    add_img_fit(s, img("mosaic"), 0.45 + (_bw - _w) / 2, 1.22, w=_w, name="Mosaic DAI")
    text(s, 0.45, 4.92, 7.0, 0.3, cap("mosaic", "Ghép ảnh: nhiều gương mặt Việt Nam — nguồn ảnh ở phụ lục"),
         size=9.5, color="cream", italic=True, align="c", name="Caption")
    add_img(s, img("s5strip"), 0.45, 5.3, 7.0, 1.2, focus=(0.42, 0.18), name="Anh tu lieu Cong giao 1945")
    text(s, 0.45, 6.54, 7.0, 0.3, "2/9/1945: khối Trường Thần học Công giáo, mít tinh Ngày Độc lập · Bảo tàng Lịch sử Quốc gia (phạm vi công cộng)",
         size=8.5, color="cream", italic=True, align="c", name="Caption 2")
    x = 7.85
    text(s, x, 0.62, 4.9, 0.4, "2a · NỘI DUNG", font=BODY_SEMI, size=13, color="gold2", char=3)
    q1 = text(s, x, 1.05, 5.0, 1.5, [[{"t": "AI NẰM TRONG", "size": 30, "color": "ivory"}],
                                     [{"t": "CHỮ “ĐẠI”?", "size": 46, "color": "gold2"}]],
              font=DISPLAY, bold=True, spacing=0.95)
    a1 = text(s, x, 2.62, 5.0, 0.8, "TOÀN DÂN", font=DISPLAY, size=44, bold=True, color="ivory", name="Answer")
    a2 = text(s, x, 3.5, 4.95, 1.3, [
        [{"t": "Chủ thể: ", "bold": True, "color": "gold2"},
         {"t": "toàn thể nhân dân — mọi người\nViệt Nam yêu nước, không phân biệt dân tộc,\ntôn giáo, đảng phái, giai cấp, tầng lớp,\ngià trẻ, gái trai, giàu nghèo."}]],
        size=14.5, color="cream", spacing=1.05)
    a3 = text(s, x, 4.6, 4.95, 0.6, [[{"t": "Nền tảng: ", "bold": True, "color": "gold2"},
                                       {"t": "công nhân, nông dân và trí thức\n(liên minh công – nông – trí)."}]], size=14.5, color="cream", spacing=1.05)
    qb = rect(s, x, 5.16, 4.95, 1.72, fill=None, line=C["gold2"], line_w=1.25, radius=6000, name="Quote box")
    qt = text(s, x + 0.22, 5.27, 4.55, 0.9, N.Q["khangchien"]["short"], font=DISPLAY, size=13, italic=True, color="ivory", spacing=1.0)
    qc = text(s, x + 0.22, 6.4, 4.55, 0.42, N.Q["khangchien"]["cite"].replace(" · Toàn tập", "\nToàn tập"), size=9.5, color="gold2", spacing=1.0)
    content_tags(s, ["A"], dark=True)
    folio(s, 5, TOTAL, dark=True)
    fade_sequence(s, [[a1, a2, a3], [qb, qt, qc]])
    transition_fade(s)
    notes(s, N.S[5])


# =====================================================================  S6 SIGNATURE
def s06():
    s = blank(prs)
    add_bg(s, cre("bg_burgundy_center.png"))
    text(s, 0.4, 0.9, W - 0.8, 1.6, [[{"t": "ĐOÀN KẾT ", "color": "ivory"}, {"t": "≠", "color": "gold2", "size": 80},
                                       {"t": " ĐỒNG NHẤT", "color": "ivory"}]],
         font=DISPLAY, size=66, bold=True, align="c", anchor="m", name="Signature")
    # left: identical
    lx, ly = 1.0, 3.35
    for i in range(7):
        rect(s, lx + i * 0.62, ly + 1.05, 0.46, 0.46, fill=C["cream"], alpha=0.55, shape=MSO_SHAPE.OVAL, name="Same")
    text(s, lx, ly + 0.1, 4.4, 0.9, [[{"t": "ĐỒNG NHẤT", "bold": True, "color": "cream", "char": 3, "size": 13}],
                                      [{"t": "mọi người phải giống hệt nhau", "font": DISPLAY, "italic": True, "size": 19, "color": "cream"}]],
         spacing=1.1)
    line(s, 6.66, 3.3, 6.66, 5.2, C["gold2"], width=0.75, alpha=0.6)
    add_img_fit(s, cre("convergence.png"), 7.2, 2.75, w=3.4, name="Hoi tu")
    text(s, 10.55, ly + 0.1, 2.4, 1.6, [[{"t": "ĐOÀN KẾT", "bold": True, "color": "gold2", "char": 3, "size": 13}],
                                         [{"t": "khác nhau — nhưng\ncùng hướng về\nđiểm chung", "font": DISPLAY, "italic": True,
                                           "size": 19, "color": "ivory"}]], spacing=1.1)
    qq = text(s, 0.7, 5.55, W - 1.4, 0.55, "Nếu con người vốn khác nhau, điều gì khiến họ vẫn có thể cùng đứng trong một khối đoàn kết?",
              font=DISPLAY, size=19, italic=True, color="gold2", align="c", name="Question")
    text(s, 0.7, 6.2, W - 1.4, 0.68, [[{"t": N.Q["tuykhac"]["short"], "italic": True, "bold": True, "color": "ivory", "size": 13},
                                       {"t": "  — Hồ Chí Minh", "color": "cream", "size": 12}],
                                      [{"t": "nói với cán bộ, công nhân Nhà máy điện Yên Phụ và Nhà máy đèn Bờ Hồ, 12-1954 · Toàn tập, t.9, tr.203", "color": "cream", "size": 9.5}]],
         align="c", spacing=1.05, name="HCM tuy khac")
    content_tags(s, ["A", "B"], dark=True)
    folio(s, 6, TOTAL, dark=True)
    transition_fade(s, "slow")
    notes(s, N.S[6])


# =====================================================================  S7 SYSTEM MODEL (HAND)
def s07():
    s = blank(prs)
    add_bg(s, cre("paper_cream.png"))
    text(s, 0.6, 0.5, 8.0, 0.4, "2b · ĐIỀU KIỆN  |  3 · HÌNH THỨC TỔ CHỨC – NGUYÊN TẮC", font=BODY_SEMI, size=13, color="red", char=2.5)
    text(s, 0.6, 0.9, 7.6, 1.2, ["ĐIỀU GÌ GIỮ NHỮNG KHÁC BIỆT", "Ở CÙNG NHAU?"], font=DISPLAY, size=30, bold=True,
         color="burg", spacing=0.98)
    fingers = [  # x, top, title, sub, kind
        (0.62, 2.55, "MẶT TRẬN\nDÂN TỘC\nTHỐNG NHẤT", "nền tảng\ncông – nông – trí;\ndo Đảng lãnh đạo", "nguyên tắc 1"),
        (2.02, 2.3, "TRUYỀN\nTHỐNG", "yêu nước –\nnhân nghĩa –\nđoàn kết", "điều kiện 1"),
        (3.42, 2.08, "KHOAN DUNG,\nĐỘ LƯỢNG", "với con người", "điều kiện 2"),
        (4.82, 2.3, "NIỀM TIN\nVÀO\nNHÂN DÂN", "", "điều kiện 3"),
        (6.22, 2.8, "HIỆP THƯƠNG\nDÂN CHỦ", "bàn bạc để\nđi đến nhất trí", "nguyên tắc 3"),
    ]
    fw = 1.32
    cols = [C["red"], C["crim"], C["burg"], C["crim"], C["red"]]
    for i, (fx, top, t, sub, kind) in enumerate(fingers):
        f = rect(s, fx, top, fw, 5.9 - top, fill=cols[i], radius=50000, name=f"Finger {i+1}")
        text(s, fx + 0.1, top + 0.3, fw - 0.2, 0.3, kind.upper(), size=8.5, bold=True, color="gold2", align="c", char=1)
        text(s, fx + 0.05, top + 0.58, fw - 0.1, 1.9, [[{"t": t, "bold": True, "size": 11.5, "color": "ivory"}],
                                                       [{"t": sub, "size": 10.5, "color": "cream"}]], align="c", spacing=1.0, para_after=3)
    palm = rect(s, 0.5, 4.8, 7.2, 1.62, fill=C["burg"], radius=30000, name="Palm")
    shadow(palm, blur=16, dist=4, alpha=0.22)
    text(s, 0.95, 4.86, 6.35, 1.5, [
        [{"t": "NGUYÊN TẮC 2 CỦA MẶT TRẬN", "size": 10.5, "bold": True, "color": "gold2", "char": 2.5}],
        [{"t": "BẢO ĐẢM LỢI ÍCH TỐI CAO CỦA DÂN TỘC", "font": DISPLAY, "size": 21, "bold": True, "color": "ivory"}],
        [{"t": "và quyền lợi cơ bản của các tầng lớp nhân dân", "font": DISPLAY, "size": 16, "italic": True, "color": "cream"}]],
        align="c", anchor="m", spacing=1.05)
    arrow(s, 7.85, 5.6, 8.55, 5.6, C["gold"], width=2.25)
    text(s, 8.7, 5.05, 4.2, 1.1, [[{"t": "SỨC MẠNH", "color": "red"}], [{"t": "CHUNG", "color": "red"}]],
         font=DISPLAY, size=30, bold=True, spacing=0.92)
    qcard = rect(s, 8.7, 1.05, 4.15, 3.5, fill=C["white"], radius=6000, name="Quote ban tay")
    shadow(qcard, blur=14, dist=3, alpha=0.12)
    add_img_fit(s, cre("hand_5_burg.png"), 8.95, 1.3, h=0.85)
    text(s, 8.95, 2.25, 3.7, 1.75, N.Q["nangontay"]["short"], font=DISPLAY, size=16, italic=True, color="burg", spacing=1.05)
    text(s, 8.95, 3.95, 3.7, 0.5, N.Q["nangontay"]["cite"].replace(" · Toàn tập", "\nToàn tập"), size=10, color="muted")
    text(s, 8.7, 6.2, 4.2, 0.7, "Mô hình “bàn tay” là cách nhóm sắp xếp các luận điểm. Khung theo bài giảng Chương 5 của lớp: 3 điều kiện, 4 nguyên tắc của Mặt trận (nguyên tắc 4: đoàn kết chặt chẽ, lâu dài, thật sự, chân thành).",
         size=9.5, italic=True, color="muted", spacing=1.0)
    content_tags(s, ["A", "B"], dark=False)
    folio(s, 7, TOTAL, dark=False)
    transition_fade(s)
    notes(s, N.S[7])


# =====================================================================  S8 CLASSROOM EXPERIMENT
def s08():
    s = blank(prs)
    add_bg(s, cre("paper_cream_soft.png"))
    text(s, 0.6, 0.5, 6.0, 0.4, "THÍ NGHIỆM TRONG LỚP", font=BODY_SEMI, size=13, color="red", char=3)
    text(s, 0.6, 0.88, 7.6, 0.9, "KHI BỐN NHÓM ĐỀU CÓ LÝ…", font=DISPLAY, size=36, bold=True, color="burg")
    text(s, 8.2, 0.84, 4.65, 0.62, "20.000.000 VNĐ", font=BODY_SEMI, size=31, color="red", align="r")
    text(s, 8.2, 1.42, 4.65, 0.36, "quỹ lớp (giả định) · chỉ đủ cho 1 dự án cộng đồng", size=13, color="muted", align="r")
    groups = [("s8_1", "NHÓM 1 · TRẺ EM", "Hỗ trợ trẻ em khó khăn —\ncác em cần nhất."),
              ("s8_2", "NHÓM 2 · MÔI TRƯỜNG", "Làm sạch, trồng cây —\nlợi cho tất cả."),
              ("s8_3", "NHÓM 3 · NGƯỜI CAO TUỔI", "Thăm và chăm sóc\nngười cao tuổi neo đơn."),
              ("s8_4", "NHÓM 4 · HỌC LIỆU", "Sách, tài liệu cho những\nbạn thiếu điều kiện.")]
    x0, y0, cw, chh, g = 0.6, 1.95, 2.92, 2.6, 0.19
    for i, (k, t, d) in enumerate(groups):
        cx = x0 + i * (cw + g)
        add_img(s, img(k), cx, y0, cw, chh, focus=(0.5, 0.4), shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=5000, name=f"Anh nhom {i+1}")
        grad_rect(s, cx, y0 + chh - 1.75, cw, 1.75, [(0, C["deep"], 0), (0.38, C["deep"], 0.78), (1, C["deep"], 0.94)], angle=90)
        text(s, cx + 0.18, y0 + chh - 1.0, cw - 0.36, 0.3, t, size=10.5, bold=True, color="gold2", char=1.2)
        text(s, cx + 0.18, y0 + chh - 0.7, cw - 0.36, 0.66, d, size=12.5, color="ivory", spacing=1.0)
    text(s, 0.6, 4.64, 8.4, 0.5, [[{"t": "Nếu cả bốn đều có lý — ", "color": "burg"},
                                   {"t": "đoàn kết có phải là chọn một nhóm thắng?", "color": "red", "bold": True}]],
         font=DISPLAY, size=18, italic=True)
    text(s, 8.85, 4.72, 4.0, 0.36, "30 giây: mỗi dãy bàn chọn 1 nhóm, nêu 1 lý do.", size=11.5, color="brown", align="r", spacing=1.0)
    # reveal (click) — sits below the photos so the static PDF keeps everything visible
    rv = rect(s, 0.6, 5.22, 12.13, 1.62, fill=C["burg"], radius=6000, name="Reveal card")
    shadow(rv, blur=18, dist=4, alpha=0.3)
    r1 = text(s, 0.9, 5.32, 5.6, 0.3, "GỢI Ý CỦA NHÓM: MỤC TIÊU Ở TẦM CAO HƠN", size=10.5, bold=True, color="gold2", char=1.5)
    r2 = text(s, 0.9, 5.6, 5.6, 0.55, [[{"t": "“Góc đọc xanh liên thế hệ” "}, {"t": "(ví dụ)", "size": 13, "bold": False, "italic": True, "color": "cream"}]], font=DISPLAY, size=23, bold=True, color="ivory")
    r3 = text(s, 0.9, 6.13, 5.6, 0.66, "Ông bà neo đơn đọc sách cùng các em nhỏ khó khăn · kệ sách tái chế · bồn cây do lớp trồng · học liệu cho bạn thiếu điều kiện.",
              size=11.5, color="cream", spacing=1.0)
    r5 = line(s, 6.62, 5.42, 6.62, 6.64, C["gold2"], width=0.75, alpha=0.6)
    r4 = text(s, 6.85, 5.32, 5.7, 1.48, [
        [{"t": "Liên hệ lý luận: ", "bold": True, "color": "gold2"},
         {"t": "thống nhất mục tiêu và lợi ích, giải quyết hài hòa lợi ích chung và lợi ích riêng; bàn bạc công khai theo tinh thần hiệp thương dân chủ để đi đến nhất trí."}],
        [{"t": "Hồ Chí Minh bàn về cách lãnh đạo: " + N.Q["sosanh"]["short"] + " ", "italic": True, "size": 10.5, "color": "cream"},
         {"t": "(Sửa đổi lối làm việc, 10-1947 · Toàn tập, t.5, tr.336)", "size": 9, "color": "gold2"}]],
        size=12.5, color="ivory", spacing=1.02, para_after=3)
    content_tags(s, ["B", "C"], dark=False)
    folio(s, 8, TOTAL, dark=False)
    fade_sequence(s, [[rv, r1, r2, r3, r4, r5]])
    transition_fade(s)
    notes(s, N.S[8])


# =====================================================================  S9 SCALE UP
def s09():
    s = blank(prs)
    add_bg(s, cre("bg_burgundy.png"))
    text(s, 0.6, 0.5, 6.0, 0.4, "PHÓNG TO QUY MÔ", font=BODY_SEMI, size=13, color="gold2", char=3)
    text(s, 0.6, 0.85, 9.0, 0.8, "Từ một lớp học đến cả quốc gia", font=DISPLAY, size=32, bold=True, color="ivory")
    panels = [("s9_1", "LỚP HỌC", "sở thích, cách làm việc", 2.15),
              ("s9_2", "TRƯỜNG", "ngành học, quê quán, vùng miền", 2.6),
              ("s9_3", "CỘNG ĐỒNG", "thế hệ, nghề nghiệp, tín ngưỡng", 3.05),
              ("s9_4", "QUỐC GIA", "dân tộc, tôn giáo, giai tầng, lợi ích", 3.5)]
    x, y, hh = 0.6, 1.8, 2.45
    for i, (k, t, d, w) in enumerate(panels):
        add_img(s, img(k), x, y + (3 - i) * 0.25, w, hh - (3 - i) * 0.25, focus=(0.5, 0.45), name=f"Quy mo {i+1}")
        text(s, x, y + hh + 0.08, w, 0.4, t, font=DISPLAY, size=18, bold=True, color="gold2")
        text(s, x, y + hh + 0.45, w, 0.6, d, size=12, color="cream", spacing=1.0)
        if i < 3:
            arrow(s, x + w + 0.07, y + hh - 0.6, x + w + 0.2, y + hh - 0.6, C["gold2"], width=2.5)
        x += w + 0.27
    # analysis + evidence
    text(s, 0.6, 5.3, 5.4, 1.62, [
        [{"t": "Quy mô càng lớn, khác biệt càng phức tạp.", "bold": True, "color": "ivory"}],
        [{"t": "Vì vậy càng cần điểm tương đồng ở tầm cao hơn và một hình thức\ntổ chức đủ rộng: Mặt trận dân tộc thống nhất do Đảng lãnh đạo.", "color": "cream"}],
        [{"t": "NQ 23-NQ/TW (2003): “…lấy mục tiêu giữ vững độc lập, thống nhất của Tổ quốc, vì dân giàu, nước mạnh, xã hội công bằng, dân chủ, văn minh làm điểm tương đồng…”", "italic": True, "size": 10.5, "color": "gold2"}]],
        size=12.5, spacing=1.05, para_after=3)
    ev = N.EVIDENCE
    ex, bw = 6.15, 2.12
    for i, e in enumerate(ev[:3]):
        bx = ex + i * (bw + 0.11)
        rect(s, bx, 5.42, bw, 1.42, fill=C["ivory"], alpha=0.07, line=C["gold2"], line_w=0.75, line_alpha=0.6, radius=6000)
        text(s, bx + 0.15, 5.5, bw - 0.3, 0.5, e["big"], font=BODY_SEMI, size=19, color="gold2")
        text(s, bx + 0.15, 5.98, bw - 0.3, 0.84, e["label"], size=10, color="cream", spacing=1.0)
    content_tags(s, ["B", "C"], dark=True)
    folio(s, 9, TOTAL, dark=True)
    source_note(s, N.EVIDENCE_SRC, dark=True, w=8.0, y=6.95)
    transition_fade(s)
    notes(s, N.S[9])


# =====================================================================  S10 FEED
def s10():
    s = blank(prs)
    add_bg(s, cre("paper_cream.png"))
    text(s, 0.6, 0.5, 6.4, 0.4, "THÍ NGHIỆM TƯ DUY", font=BODY_SEMI, size=13, color="red", char=3)
    text(s, 0.6, 0.88, 6.6, 1.6, ["ĐẠI ĐOÀN KẾT", "KHI MỖI NGƯỜI MỘT “FEED”?"], font=DISPLAY, size=29, bold=True,
         color="burg", spacing=0.98)
    text(s, 0.6, 2.2, 6.65, 1.3, "Nếu mỗi người tiếp nhận một luồng thông tin khác nhau,\nviệc nhận ra “điểm chung” có khó hơn không?",
         font=DISPLAY, size=19, italic=True, color="brown", spacing=1.05)
    habits = [("Lắng nghe trước khi phản hồi", "khoan dung, độ lượng với con người"),
              ("Tranh luận có văn hóa, để tìm điểm chung", "tinh thần hiệp thương dân chủ: bàn bạc để đi đến nhất trí"),
              ("Nhắc lại mục tiêu chung", "thống nhất mục tiêu và lợi ích")]
    y = 3.62
    for i, (h, p) in enumerate(habits):
        rect(s, 0.6, y, 0.46, 0.46, fill=C["burg"], shape=MSO_SHAPE.OVAL)
        text(s, 0.6, y, 0.46, 0.46, str(i + 1), font=BODY_SEMI, size=16, color="gold2", align="c", anchor="m")
        text(s, 1.25, y - 0.04, 5.6, 0.75, [[{"t": h, "bold": True, "size": 16.5, "color": "ink"}],
                                           [{"t": "↔ " + p, "size": 12.5, "color": "muted"}]], spacing=1.0)
        y += 0.82
    text(s, 0.6, 6.12, 6.3, 0.7, N.FEED_STAT, size=10.5, color="brown", spacing=1.0, name="Stat")
    # phones
    palettes = [[C["red"], C["warm"], "E6A49A", C["crim"]], [C["gold"], C["gold2"], "E8D3A6", C["brown"]],
                [C["brown"], "8C6F62", "C9B3A6", C["muted"]]]
    for i, pal in enumerate(palettes):
        px, py, pw, ph = 7.35 + i * 1.88, 1.0 + (i % 2) * 0.3, 1.68, 3.1
        body = rect(s, px, py, pw, ph, fill="1A1213", radius=12000, name=f"Phone {i+1}")
        shadow(body, blur=18, dist=5, alpha=0.3)
        rect(s, px + 0.09, py + 0.12, pw - 0.18, ph - 0.24, fill=C["ivory"], radius=9000)
        yy = py + 0.32
        for j in range(5):
            hh = 0.36 if j % 2 == 0 else 0.62
            rect(s, px + 0.19, yy, pw - 0.38, hh, fill=pal[j % 4], radius=12000)
            yy += hh + 0.09
        text(s, px, py + ph + 0.1, pw, 0.3, f"Feed {'ABC'[i]}", size=11, bold=True, color="muted", align="c", char=1.5)
    text(s, 7.35, 4.98, 5.5, 1.0, N.FEED_QUOTE["short"].replace("trách nhiệm, trung", "trách nhiệm,\ntrung").replace("tin giả, không gieo", "tin giả,\nkhông gieo").replace("hiểu biết, tinh", "hiểu biết,\ntinh"), font=DISPLAY, size=12.5, italic=True, color="burg", spacing=1.03)
    text(s, 7.35, 5.98, 5.5, 0.3, N.FEED_QUOTE["cite"], size=9.5, color="muted")
    text(s, 7.35, 6.44, 5.5, 0.45, "Thí nghiệm tư duy của nhóm — không khẳng định thuật toán gây ra chia rẽ.",
         size=10.5, italic=True, color="muted")
    content_tags(s, ["B", "C"], dark=False)
    folio(s, 10, TOTAL, dark=False)
    transition_fade(s)
    notes(s, N.S[10])


# =====================================================================  S11 UNITY LAB
def s11():
    s = blank(prs)
    add_bg(s, cre("bg_burgundy_center.png"))
    add_img(s, cre("drum_gold_14.png"), 4.65, 0.45, 6.6, 6.6, name="Motif")
    text(s, 0.6, 0.5, 5.0, 0.4, "TRẢI NGHIỆM BẤT NGỜ", font=BODY_SEMI, size=13, color="gold2", char=4)
    text(s, 0.6, 1.05, 5.4, 2.6, [[{"t": "UNITY", "color": "ivory"}], [{"t": "LAB", "color": "gold2"}]],
         font=DISPLAY, size=82, bold=True, spacing=0.8, name="Wordmark")
    text(s, 0.6, 3.55, 4.9, 0.95, "Bạn xử lý khác biệt như thế nào?", font=DISPLAY, size=24, italic=True, color="ivory", spacing=1.0)
    text(s, 0.6, 4.58, 4.9, 0.4, "3 tình huống  ·  ~2 phút  ·  không chấm đúng / sai", size=14, color="gold2", bold=True)
    cta = rect(s, 0.6, 5.25, 3.6, 0.72, fill=C["gold2"], radius=50000, name="CTA")
    text(s, 0.6, 5.25, 3.6, 0.72, "QUÉT MÃ BÊN PHẢI  →", size=17, bold=True, color="burg", align="c", anchor="m", char=1.5)
    text(s, 0.6, 6.1, 5.0, 0.6, "Cuối game, mỗi bạn nhận một “thẻ màu” —\ncả lớp giơ tay theo màu.", size=12, color="cream")
    add_img_fit(s, img("s11_phone"), 5.6, 0.42, h=6.65, name="Phone mockup")
    qr_png, url_file = cre("qr_unitylab.png"), ROOT / "_build" / "unitylab_url.txt"
    if qr_png.exists() and url_file.exists():  # real QR (written by _build/make_qr.py after deployment)
        url = url_file.read_text(encoding="utf-8").strip()
        s.shapes.add_picture(str(qr_png), Inches(9.75), Inches(1.55), Inches(2.95), Inches(2.95)).name = "QR_UNITYLAB"
        caption = url.split("://", 1)[-1].rstrip("/")
    else:  # no live URL yet: honest placeholder, never a fake QR
        rect(s, 9.75, 1.55, 2.95, 2.95, fill=C["ivory"], line=C["gold2"], line_w=2, radius=6000, dash="dash", name="QR_PLACEHOLDER")
        text(s, 9.95, 2.2, 2.55, 1.6, [[{"t": "QR PLACEHOLDER", "bold": True, "size": 15, "color": "burg"}],
                                       [{"t": "UPDATE AFTER DEPLOYMENT", "size": 11, "color": "red", "char": 1}]],
             align="c", anchor="m", spacing=1.1, name="QR_PLACEHOLDER_TEXT")
        caption = "QR PLACEHOLDER — UPDATE AFTER DEPLOYMENT"
    text(s, 9.75, 4.62, 2.95, 0.4, "Mở camera điện thoại để quét", size=13, color="gold2", align="c")
    text(s, 9.5, 5.02, 3.45, 0.5, caption, size=11 if "PLACEHOLDER" not in caption else 9.5, color="cream", align="c", name="URL caption")
    content_tags(s, ["C"], dark=True)
    folio(s, 11, TOTAL, dark=True)
    transition_fade(s, "slow")
    notes(s, N.S[11])


# =====================================================================  S12 RETURN
def s12():
    s = blank(prs)
    add_bg(s, cre("paper_cream.png"))
    add_img(s, img("s12"), 8.75, 0, W - 8.75, H, focus=(0.5, 0.3), name="Chan dung")
    grad_rect(s, 8.75, 3.3, W - 8.75, 4.2, [(0, C["deep"], 0), (0.45, C["deep"], 0.8), (1, C["deep"], 0.96)], angle=90)
    text(s, 9.05, 4.62, 4.0, 1.55, N.Q["doanket3"]["short"].replace("đoàn kết, đại đoàn kết, ", "đoàn kết,\nđại đoàn kết,\n").replace("thành công, đại", "thành công,\nđại"),
         font=DISPLAY, size=17, italic=True, bold=True, color="gold2", spacing=1.0)
    text(s, 9.05, 6.2, 4.0, 0.62, N.Q["doanket3"]["cite"], size=9.5, color="cream", spacing=1.0)
    text(s, 0.6, 0.5, 7.6, 0.4, "QUAY LẠI CÂU HỎI BAN ĐẦU", font=BODY_SEMI, size=13, color="red", char=3)
    text(s, 0.6, 0.86, 7.9, 0.5, "Đoàn kết có phải là tất cả mọi người phải giống nhau?", font=DISPLAY, size=20,
         italic=True, color="muted")
    a = text(s, 0.6, 1.4, 8.0, 0.6, "KHÔNG PHẢI GIỐNG NHAU VỀ MỌI MẶT.", font=DISPLAY, size=25, bold=True,
             color="burg", name="Reveal 1")
    # two columns
    cw = 3.85
    c1 = rect(s, 0.6, 2.1, cw, 2.3, fill=C["burg"], radius=6000, name="Can nhat tri")
    t1 = text(s, 0.85, 2.18, cw - 0.45, 2.14, anchor="m", paras=[
        [{"t": "CẦN NHẤT TRÍ", "size": 12, "bold": True, "color": "gold2", "char": 2}],
        [{"t": "Mục đích và lập trường.", "font": DISPLAY, "size": 17, "bold": True, "color": "ivory"}],
        [{"t": "“Đoàn kết thực sự nghĩa là mục đích phải nhất trí và lập trường cũng phải nhất trí.”", "size": 12.5, "italic": True, "color": "cream"}],
        [{"t": "Hồ Chí Minh, 19-3-1958 · Toàn tập, t.11, tr.362", "size": 9, "color": "gold2"}]], spacing=1.03, para_after=4)
    c2 = rect(s, 0.6 + cw + 0.2, 2.1, cw, 2.3, fill=C["white"], radius=6000, name="Co the khac")
    shadow(c2, blur=12, dist=3, alpha=0.1)
    t2 = text(s, 0.85 + cw + 0.2, 2.18, cw - 0.45, 2.14, anchor="m", paras=[
        [{"t": "CÓ THỂ KHÁC NHAU", "size": 12, "bold": True, "color": "red", "char": 2}],
        [{"t": "Dân tộc, tôn giáo, tuổi tác, nghề nghiệp, cá tính — kể cả quá khứ, nếu nay thật thà tán thành mục tiêu chung.", "font": BODY, "size": 14, "bold": True, "color": "burg"}],  # Phudu is caps-only: sentences use the body face
        [{"t": "“Trong mấy triệu người cũng có người thế này thế khác…” · “không chia tôn giáo, đảng phái, dân tộc” (1946)", "size": 10.5, "italic": True, "color": "brown"}],
        [{"t": "Toàn tập, t.4, tr.280 & 534 · t.9, tr.203 & 244", "size": 9, "color": "red"}]], spacing=1.03, para_after=3)
    b = text(s, 0.6, 4.66, 7.9, 0.66, [[{"t": "→ ", "color": "gold"}, {"t": "Đoàn kết là ", "color": "ink"},
                                         {"t": "nhất trí về mục đích, lập trường và cùng hướng tới điểm chung", "color": "red", "bold": True},
                                         {"t": "; tôn trọng khác biệt chính đáng — nhưng không vô nguyên tắc.", "color": "ink"}]],
             font=BODY, size=15.5, spacing=1.0, name="Reveal 3")
    btag = pill(s, 0.6, 5.42, "DIỄN GIẢI CỦA NHÓM", None, "burg", size=9.5, line=C["burg"], w=2.05)
    poll = text(s, 0.6, 5.86, 7.9, 0.95, [
        [{"t": "Nhìn lại 4 đáp án: ", "bold": True, "color": "burg"},
         {"t": "B gần nhất · A đúng một phần (nhất trí mục đích, lập trường — không phải mọi ý kiến) · "
               "C là cách ra quyết định hợp lệ, chưa phải định nghĩa đoàn kết · D trái với “vừa đoàn kết, vừa đấu tranh” (t.11, tr.362)."}]],
        size=11.5, color="ink", spacing=1.03, name="Poll recap")
    fade_sequence(s, [[a], [c1, t1, c2, t2], [b, btag], [poll]])
    content_tags(s, ["A"], dark=True)
    folio(s, 12, TOTAL, dark=False)
    transition_fade(s)
    notes(s, N.S[12])


# =====================================================================  S13 AI USAGE
def s13():
    s = blank(prs)
    add_bg(s, cre("paper_cream_soft.png"))
    text(s, 0.6, 0.5, 8.0, 0.4, "AI USAGE  ·  LIÊM CHÍNH HỌC THUẬT", font=BODY_SEMI, size=13, color="red", char=3)
    text(s, 0.6, 0.88, 12.0, 0.8, "AI hỗ trợ gì — và nhóm chịu trách nhiệm gì?", font=DISPLAY, size=30, bold=True, color="burg")
    colw = 5.95
    for i, (title, sub, items, fill, tcol) in enumerate([
        ("AI HỖ TRỢ", N.AI_TOOL, N.AI_DID, C["burg"], "ivory"),
        ("SINH VIÊN CHỊU TRÁCH NHIỆM", "Toàn bộ thành viên nhóm", N.SV_DID, C["white"], "ink")]):
        x = 0.6 + i * (colw + 0.23)
        card = rect(s, x, 1.85, colw, 3.55, fill=fill, radius=5000)
        shadow(card, blur=12, dist=3, alpha=0.12)
        text(s, x + 0.3, 2.0, colw - 0.6, 0.4, title, size=13, bold=True, color="gold2" if i == 0 else "red", char=2.5)
        text(s, x + 0.3, 2.36, colw - 0.6, 0.4, sub, size=11, italic=True, color="cream" if i == 0 else "muted")
        tb = text(s, x + 0.3, 2.8, colw - 0.6, 2.0, [[{"t": "•\t" + it}] for it in items], size=13.5,
                  color=tcol, spacing=1.0, para_after=3, hang=0.24)
        if i == 0:
            text(s, x + 0.3, 4.62, colw - 0.6, 0.72, [[{"t": "Ví dụ AI tự sai: ", "bold": True, "color": "gold2"},
                {"t": "vòng tìm nguồn ghi một câu là “không có trong Toàn tập”; vòng phản biện tra lại và tìm thấy ở t.14, tr.27 → đã sửa."}]],
                size=11, color="cream", spacing=1.0, name="AI error example")
        else:
            text(s, x + 0.3, 4.62, colw - 0.6, 0.72, [[{"t": "Nhật ký chỉnh sửa của nhóm: ", "bold": True, "color": "red"},
                {"t": "07_AI-USAGE.md, mục 5 (việc gì, ai làm, sửa gì so với bản AI)."}]],
                size=11, color="muted", spacing=1.0, name="Student edit log")
    box = rect(s, 0.6, 5.62, 12.13, 1.2, fill=C["paper2"], radius=5000, name="Cam ket")
    text(s, 0.85, 5.66, 11.7, 1.12, [[{"t": "Cam kết: ", "bold": True, "color": "burg"}, {"t": N.INTEGRITY}],
                                     [{"t": "Bản cam kết có chữ ký và chi tiết công cụ, prompt, kết quả AI: 07_AI-USAGE.md (mục 1–5, 9)", "size": 10.5, "color": "muted"}]],
         size=12, color="ink", spacing=1.05, para_after=3, anchor="m")
    folio(s, 13, TOTAL, dark=False)
    transition_fade(s)
    notes(s, N.S[13])


# =====================================================================  S14 SOURCES (appendix)
def s14():
    s = blank(prs)
    add_bg(s, cre("paper_cream_soft.png"))
    text(s, 0.6, 0.5, 8.0, 0.4, "PHỤ LỤC", font=BODY_SEMI, size=13, color="red", char=3)
    text(s, 0.6, 0.85, 12.0, 0.7, "Nguồn tham khảo & ảnh", font=DISPLAY, size=28, bold=True, color="burg")
    colw = 5.95
    for i, (title, items) in enumerate(N.SOURCES_SLIDE):
        x = 0.6 + (i % 2) * (colw + 0.23)
        y = 1.7 + (i // 2) * 2.3
        text(s, x, y, colw, 0.35, title, size=11.5, bold=True, color="red", char=2)
        text(s, x, y + 0.38, colw, 2.2, [[{"t": it}] for it in items], size=10, color="ink", spacing=1.0, para_after=2)
    folio(s, 14, TOTAL, dark=False)
    transition_fade(s)
    notes(s, N.S[14])


# =====================================================================  S15 CLOSING
def s15():
    s = blank(prs)
    add_bg(s, cre("bg_burgundy_center.png"))
    add_img(s, img("s15"), 0, 0, W, H, focus=(0.5, 0.45), alpha=0.28, name="Anh nen")
    band(s, 1.3)
    add_img_fit(s, cre("lotus_gold.png"), W / 2 - 0.75, 0.55, w=1.5)
    text(s, 0.5, 1.75, W - 1.0, 1.2, "XIN CẢM ƠN", font=DISPLAY, size=66, bold=True, color="ivory", align="c", char=4)
    text(s, 0.5, 2.95, W - 1.0, 0.55, "Mời các nhóm phản biện", font=DISPLAY, size=24, italic=True, color="gold2", align="c")
    qs = N.DISCUSSION
    for i, q in enumerate(qs):
        x = 0.9 + i * 3.92
        rect(s, x, 3.85, 3.7, 1.75, fill=C["deep"], alpha=0.55, line=C["gold2"], line_w=0.75, line_alpha=0.7, radius=6000)
        text(s, x + 0.22, 3.95, 3.3, 0.4, f"CÂU HỎI {i+1}", size=10.5, bold=True, color="gold2", char=2)
        text(s, x + 0.22, 4.3, 3.3, 1.25, q, size=13.5, color="ivory", spacing=1.05)
    text(s, 0.5, 5.7, W - 1.0, 0.36, cap("s15", ""), size=9, italic=True, color="cream", align="c", name="Caption")
    transition_fade(s, "slow")
    notes(s, N.S[15])


for fn in (s01, s02, s03, s04, s05, s06, s07, s08, s09, s10, s11, s12, s13, s14, s15):
    fn()

prs.core_properties.title = "Tư tưởng Hồ Chí Minh về đại đoàn kết dân tộc"
prs.core_properties.subject = "Học phần Tư tưởng Hồ Chí Minh — Chương 5"
prs.core_properties.author = "Nhóm ___ (Đại học FPT)"
prs.core_properties.keywords = "đại đoàn kết; Hồ Chí Minh; HCM202"
prs.save(str(OUT))
print("saved", OUT, len(prs.slides._sldIdLst), "slides")
