"""Generate a REAL QR code for the deployed Unity Lab and put it into the deck.

Run this only after Unity Lab is online (see 03_UnityLab/README.md):

    pip install segno
    python _build/make_qr.py https://<your-deployed-url>/

What it does
  1. writes 12_assets/created/qr_unitylab.png (dark burgundy on ivory, high error correction,
     4-module quiet zone) and saves the URL to _build/unitylab_url.txt
  2. opens 01_Tu_tuong_HCM_Dai_doan_ket.pptx, finds every shape named "QR_PLACEHOLDER"
     (slide 11), replaces it with the QR image at the same position/size,
     and replaces the text "QR PLACEHOLDER — UPDATE AFTER DEPLOYMENT" in the URL caption.
  3. saves the deck in place (a backup copy *.before-qr.pptx is kept).
Then re-export the PDF (PowerPoint: File > Export > PDF, or _build/render.ps1).
Because the PNG and the URL file are kept, later runs of build_deck.py draw the real QR by themselves.
"""
import shutil
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
DECK = ROOT / "01_Tu_tuong_HCM_Dai_doan_ket.pptx"
QR_PNG = ROOT / "12_assets" / "created" / "qr_unitylab.png"
URL_FILE = ROOT / "_build" / "unitylab_url.txt"


def display_url(url):
    """Short form printed under the QR: no scheme, no trailing slash."""
    return url.split("://", 1)[-1].rstrip("/")


def inject(deck_path, qr_png, url):
    """Replace the QR placeholder on the deck with qr_png; returns number of placeholders replaced."""
    from pptx import Presentation
    prs = Presentation(str(deck_path))
    replaced = 0
    for idx, slide in enumerate(prs.slides, 1):
        for shp in list(slide.shapes):
            if shp.name == "QR_PLACEHOLDER":
                x, y, w, h = shp.left, shp.top, shp.width, shp.height
                side = min(w, h)
                pic = slide.shapes.add_picture(str(qr_png), x + (w - side) // 2, y + (h - side) // 2, side, side)
                pic.name = "QR_UNITYLAB"
                shp._element.getparent().remove(shp._element)
                replaced += 1
                print(f"slide {idx}: placeholder replaced")
            elif shp.name == "QR_PLACEHOLDER_TEXT":
                shp._element.getparent().remove(shp._element)
            elif shp.has_text_frame and "QR PLACEHOLDER" in shp.text_frame.text:
                runs = [r for p in shp.text_frame.paragraphs for r in p.runs]
                if runs:
                    runs[0].text = display_url(url)
                    for r in runs[1:]:
                        r.text = ""
    prs.save(str(deck_path))
    return replaced


def main():
    if len(sys.argv) < 2 or not sys.argv[1].startswith(("http://", "https://")):
        sys.exit("Usage: python _build/make_qr.py https://your-deployed-unity-lab-url/")
    url = sys.argv[1].strip()
    try:
        import segno
    except ImportError:
        sys.exit("Missing dependency: run  pip install segno  and try again.")
    qr = segno.make(url, error="h")
    qr.save(str(QR_PNG), scale=24, border=4, dark="#5B0808", light="#FAF5E8")
    URL_FILE.write_text(url + "\n", encoding="utf-8")
    print("QR written:", QR_PNG, "| URL saved:", URL_FILE.name)
    backup = ROOT / "_build" / (DECK.stem + ".before-qr.pptx")
    shutil.copy2(DECK, backup)
    n = inject(DECK, QR_PNG, url)
    print(f"{n} placeholder(s) replaced. Backup: {backup.name}. Now re-export the PDF.")


if __name__ == "__main__":
    main()
